"""Lecture du texte des documents Excel, Word et PDF (lecture seule : rien n'est modifié)."""

from pathlib import Path

# Au-delà, le texte est coupé pour ne pas noyer l'agent (et limiter la dépense).
MAX_CHARS = 60_000


def _excel(path: Path) -> str:
    from openpyxl import load_workbook

    # data_only : on lit les valeurs calculées, pas les formules.
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## Feuille « {ws.title} »")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                parts.append("\t".join("" if v is None else str(v) for v in row))
    wb.close()
    return "\n".join(parts)


def _word(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for i, table in enumerate(doc.tables, 1):
        parts.append(f"## Tableau {i}")
        for row in table.rows:
            parts.append("\t".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = [f"## Page {i}\n{(page.extract_text() or '').strip()}" for i, page in enumerate(reader.pages, 1)]
    text = "\n".join(parts)
    if not any((page.extract_text() or "").strip() for page in reader.pages[:5]):
        return (
            "Ce PDF ne contient pas de texte lisible (probablement un scan ou un plan en image). "
            "Ouvre-le avec l'outil Read (paramètre pages, par exemple \"1-5\") pour voir ses pages en images."
        )
    return text


READERS = {".xlsx": _excel, ".xlsm": _excel, ".docx": _word, ".pdf": _pdf}


def extract(path: Path) -> str:
    reader = READERS.get(path.suffix.lower())
    if reader is None:
        return f"Format {path.suffix} non pris en charge (seulement {', '.join(READERS)})."
    text = reader(path)
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + f"\n[… texte coupé : {len(text)} caractères au total]"
    return text or "(document vide)"
